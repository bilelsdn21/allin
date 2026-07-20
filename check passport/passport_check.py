"""
Passport check + ID stamper
---------------------------
What it does:
  1. Reads your Excel list (names + ID numbers).
  2. Looks at every image in the 'images' folder (the file name = client name).
  3. Tells you who is matched, who sent an image but is NOT in the list,
     and who is in the list but did NOT send an image (missing passport).
  4. Writes the client's ID number onto a copy of each image, saved in 'output'.

The originals in 'images' are NEVER changed. If you get a new Excel with the
same names but different IDs, just replace the Excel and run this again -
the output is rebuilt from the originals every time.

Run it by double-clicking run.bat, or:  python passport_check.py
"""

import csv
import datetime
import json
import re
import shutil
import sys
import traceback
import unicodedata
from difflib import SequenceMatcher, get_close_matches
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, UnidentifiedImageError

try:
    import pandas as pd
except ImportError:
    pd = None
    import openpyxl

try:
    import fitz  # PyMuPDF, for reading PDF files
except ImportError:
    fitz = None

__version__ = "1.1.0"

HERE = Path(__file__).parent
IMG_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
PDF_EXTS = {".pdf"}
ALL_EXTS = IMG_EXTS | PDF_EXTS
PDF_DPI = 200  # quality when turning a PDF page into an image (good for printing)


# --------------------------------------------------------------------------- #
# errors the app can explain to the user (instead of crashing)
# --------------------------------------------------------------------------- #
class CheckError(Exception):
    """A problem we can explain in plain words. .message_fr holds the French."""
    def __init__(self, message, message_fr=""):
        super().__init__(message)
        self.message_fr = message_fr


class ExcelNotFound(CheckError):
    pass


class ColumnNotFound(CheckError):
    def __init__(self, wanted, headers):
        self.wanted = wanted
        self.headers = headers
        cols = ", ".join(str(h) for h in headers if str(h).strip())
        super().__init__(
            f"Could not find the '{wanted}' column in the Excel.\n"
            f"Columns found: {cols}\n"
            f"Rename the column in Excel, or set it in the Settings.",
            f"Impossible de trouver la colonne '{wanted}' dans le fichier Excel.\n"
            f"Colonnes trouvées : {cols}\n"
            f"Renommez la colonne dans Excel, ou choisissez-la dans les réglages.")


def log_error(context, exc):
    """Append a technical traceback to app_errors.log (for support/debugging)."""
    try:
        with open(HERE / "app_errors.log", "a", encoding="utf-8") as f:
            f.write(f"\n[{datetime.datetime.now():%Y-%m-%d %H:%M:%S}] {context}\n")
            f.write("".join(traceback.format_exception(exc)))
    except OSError:
        pass  # never let logging itself crash the run


# --------------------------------------------------------------------------- #
# run history + decision log (for the testing period)
# --------------------------------------------------------------------------- #
DECISION_COLUMNS = ["file", "normalized name", "decision", "matched person",
                    "id written", "index", "detail", "suggested file name"]


def write_decisions_csv(path, decisions):
    """One line per photo (and per missing person): what we decided and why.

    utf-8-sig so double-clicking the file opens correctly in Excel.
    """
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(DECISION_COLUMNS)
        for d in decisions:
            w.writerow([d.get(k, "") for k in
                        ("file", "key", "decision", "matched_name",
                         "id", "index", "detail", "suggestion")])


def archive_run(decisions, stats, cfg, report_files):
    """Copy this run's reports into runs/<timestamp>/ so nothing is
    overwritten and testing results can be compared later.

    Returns the run folder (or None if archiving failed - never fatal).
    """
    stamp = f"{datetime.datetime.now():%Y-%m-%d_%H%M%S}"
    run_dir = HERE / "runs" / stamp
    try:
        run_dir.mkdir(parents=True, exist_ok=True)
        write_decisions_csv(run_dir / "decisions.csv", decisions)
        summary = {
            "version": __version__,
            "when": stamp,
            "excel_file": str(cfg.get("excel_file", "auto")),
            "stats": stats,
        }
        with open(run_dir / "summary.json", "w", encoding="utf-8") as f:
            json.dump(summary, f, ensure_ascii=False, indent=2)
        for src in report_files:
            if src and src.exists():
                shutil.copy2(src, run_dir / src.name)
        return run_dir
    except OSError as e:
        log_error("archiving run", e)
        return None


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def load_config():
    cfg_path = HERE / "config.json"
    with open(cfg_path, "r", encoding="utf-8") as f:
        return json.load(f)


def normalize(name):
    """Make names comparable: lowercase, no accents, single spaces."""
    if name is None:
        return ""
    text = str(name).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = " ".join(text.split())
    return text


def clean_stem(stem):
    """Filename (without extension) -> comparable name.

    Also strips trailing things like '.101' in 'ANDREJA HASANOVIC.101'.
    """
    stem = re.sub(r"\.\d+$", "", stem)     # 'NAME.101' -> 'NAME'
    return normalize(stem)


# words that sometimes get typed into photo file names but are not names
JUNK_WORDS = {"passport", "passeport", "pasos", "scan", "scanned", "copy",
              "copie", "photo", "img", "image", "doc", "document", "new",
              "final", "ok"}


def diagnose_unmatched(key, fname, people, duplicates):
    """Work out WHY a photo file name matches nobody, and what to do.

    Returns a dict:
      reason      : short code ('reversed', 'junk', 'two_people',
                    'spelling', 'partial', 'unknown')
      why         : plain-English explanation + what to do
      why_fr      : the same in French
      suggestion  : the exact new file name to use (or "" if none)
    """
    all_people = {**people, **duplicates}
    all_keys = list(all_people.keys())
    key_to_name = {k: v["name"] for k, v in all_people.items()}
    ext = Path(fname).suffix or ".jpg"

    def best_ratio(a, b):
        return SequenceMatcher(None, a, b).ratio()

    # 1) name written back-to-front (surname first)?
    words = key.split()
    if len(words) >= 2:
        swapped = " ".join(reversed(words))
        if swapped in all_people:
            good = key_to_name[swapped]
            return {
                "reason": "reversed",
                "why": (f"The name is written back-to-front (family name "
                        f"first). It should be '{good}'."),
                "why_fr": (f"Le nom est écrit à l'envers (nom de famille "
                           f"d'abord). Il devrait être '{good}'."),
                "suggestion": good + ext,
            }

    # 2) extra words in the file name (e.g. 'JOHN SMITH passport scan 2')?
    stripped = " ".join(w for w in words
                        if w not in JUNK_WORDS and not w.isdigit())
    if stripped != key and stripped in all_people:
        good = key_to_name[stripped]
        return {
            "reason": "junk",
            "why": (f"The file name has extra words. It should be "
                    f"just '{good}'."),
            "why_fr": (f"Le nom du fichier contient des mots en trop. "
                       f"Il devrait être juste '{good}'."),
            "suggestion": good + ext,
        }

    # 3) two people in one file? ('A I B', 'A AND B', 'A_B', 'A & B')
    parts = re.split(r"\s+i\s+|\s+and\s+|\s+et\s+|\s*&\s*|_+", key)
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) >= 2:
        found = []
        for part in parts:
            hit = get_close_matches(part, all_keys, n=1, cutoff=0.55)
            if hit:
                found.append(key_to_name[hit[0]])
        if len(found) >= 2:
            names = " + ".join(found)
            return {
                "reason": "two_people",
                "why": (f"This file seems to hold TWO people: {names}. "
                        f"Make one copy of the file per person, each with "
                        f"one name."),
                "why_fr": (f"Ce fichier semble contenir DEUX personnes : "
                           f"{names}. Faites une copie du fichier par "
                           f"personne, chacune avec un seul nom."),
                "suggestion": "",
            }

    # 4) close spelling mistake?
    close = get_close_matches(key, all_keys, n=3, cutoff=0.6)
    if close and best_ratio(key, close[0]) >= 0.75:
        good = key_to_name[close[0]]
        return {
            "reason": "spelling",
            "why": (f"Spelling mistake in the file name. It should "
                    f"probably be '{good}'."),
            "why_fr": (f"Faute de frappe dans le nom du fichier. "
                       f"Ce devrait probablement être '{good}'."),
            "suggestion": good + ext,
        }

    # 5) part of the name matches someone (usually the family name)?
    partial_hits = []
    for w in words:
        if len(w) < 3:
            continue
        for k in all_keys:
            if w in k.split() and key_to_name[k] not in partial_hits:
                partial_hits.append(key_to_name[k])
    if partial_hits:
        show = ", ".join(partial_hits[:3])
        return {
            "reason": "partial",
            "why": (f"Part of the name matches: {show}. The rest does "
                    f"not. Check who this person really is."),
            "why_fr": (f"Une partie du nom correspond à : {show}. Le "
                       f"reste ne correspond pas. Vérifiez qui est "
                       f"vraiment cette personne."),
            "suggestion": partial_hits[0] + ext if len(partial_hits) == 1 else "",
        }

    # 6) nothing similar at all
    return {
        "reason": "unknown",
        "why": ("This name is not in the list at all. Ask if this "
                "person should be added to the list."),
        "why_fr": ("Ce nom n'est pas du tout dans la liste. Demandez "
                   "si cette personne doit être ajoutée à la liste."),
        "suggestion": "",
    }


def hex_to_rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def find_excel(setting):
    """Locate the Excel list. 'auto' = the only .xls/.xlsx in the folder."""
    if setting and setting != "auto":
        p = Path(setting)
        if not p.is_absolute():
            p = HERE / setting
        if not p.exists():
            raise ExcelNotFound(
                f"The name list was not found: {p.name}\n"
                f"Put the Excel file in and try again.",
                f"La liste des noms est introuvable : {p.name}\n"
                f"Ajoutez le fichier Excel et réessayez.")
        return p

    candidates = [p for p in HERE.iterdir()
                  if p.suffix.lower() in (".xls", ".xlsx", ".xlsm")
                  and not p.name.startswith("~$")
                  and not p.name.lower().startswith("report")]  # our own output
    if not candidates:
        raise ExcelNotFound(
            "No name list (Excel file) found.\n"
            "Put your Excel list in first, then try again.",
            "Aucune liste de noms (fichier Excel) trouvée.\n"
            "Ajoutez d'abord votre liste Excel, puis réessayez.")
    if len(candidates) > 1:
        names = ", ".join(p.name for p in candidates)
        raise ExcelNotFound(
            f"There are several Excel files: {names}\n"
            f"Keep only one, or choose one in the app.",
            f"Il y a plusieurs fichiers Excel : {names}\n"
            f"Gardez-en un seul, ou choisissez-en un dans l'application.")
    return candidates[0]


def read_excel(cfg):
    """Return (people, duplicates).

    people     : normalized_name -> {'name': original, 'id': id_str}
                 (only names that appear EXACTLY once)
    duplicates : normalized_name -> {'name': original, 'ids': [id1, id2, ...]}
                 (names that appear more than once - handled manually)
    """
    xlsx = find_excel(cfg["excel_file"])

    rows = []
    df = pd.read_excel(xlsx, dtype=str)
    headers = list(df.columns)
    rows = df.fillna("").values.tolist()

    print(f"  Excel file: {xlsx.name}")

    name_col = _pick_column(headers, cfg.get("name_column", "auto"),
                            ("name", "client", "nom", "prenom", "full"))
    id_col = _pick_column(headers, cfg.get("id_column", "auto"),
                          ("passport", "id", "cin", "number", "num", "cni"))

    # surname is optional: 'none' = names are in one column already
    surname_setting = cfg.get("surname_column", "auto")
    surname_col = None
    if str(surname_setting).lower() != "none":
        try:
            surname_col = _pick_column(headers, surname_setting,
                                       ("surname", "lastname", "last name",
                                        "family", "nom de famille"))
        except ColumnNotFound:
            surname_col = None  # no surname column, that's fine

    # index number (Excel row position) is optional
    index_col = None
    try:
        index_col = _pick_column(headers, cfg.get("index_column", "auto"),
                                 ("index", "no.", "n."))
    except ColumnNotFound:
        index_col = None

    print(f"  Name column:    '{headers[name_col]}'"
          + (f" + '{headers[surname_col]}'" if surname_col is not None else ""))
    print(f"  ID column:      '{headers[id_col]}'"
          + (f"\n  Index column:   '{headers[index_col]}'" if index_col is not None else "")
          + "\n")

    def clean_num(v):
        v = str(v).strip()
        if v.endswith(".0") and v[:-2].isdigit():
            v = v[:-2]
        return "" if v.lower() == "nan" else v

    # first pass: collect every row per normalized full name
    grouped = {}
    for r in rows:
        raw_name = str(r[name_col]).strip() if name_col < len(r) else ""
        raw_surname = (str(r[surname_col]).strip()
                       if surname_col is not None and surname_col < len(r) else "")
        raw_id = clean_num(r[id_col]) if id_col < len(r) else ""
        raw_index = (clean_num(r[index_col])
                     if index_col is not None and index_col < len(r) else "")

        full_name = (raw_name + " " + raw_surname).strip()
        if not full_name:
            continue
        key = normalize(full_name)
        grouped.setdefault(key, {"name": full_name, "ids": [], "indexes": []})
        grouped[key]["ids"].append(raw_id)
        grouped[key]["indexes"].append(raw_index)

    # split: unique names -> people ; repeated names -> duplicates (manual)
    people, duplicates = {}, {}
    for key, info in grouped.items():
        if len(info["ids"]) == 1:
            people[key] = {"name": info["name"], "id": info["ids"][0],
                           "index": info["indexes"][0]}
        else:
            duplicates[key] = info
    return people, duplicates


def _pick_column(headers, setting, keywords):
    """Find a column by config setting ('auto', a header name, or an index)."""
    norm_headers = [normalize(h) for h in headers]

    if setting != "auto":
        # exact header match?
        if normalize(setting) in norm_headers:
            return norm_headers.index(normalize(setting))
        # numeric index?
        try:
            idx = int(setting)
            if 0 <= idx < len(headers):
                return idx
        except (ValueError, TypeError):
            pass

    for i, h in enumerate(norm_headers):
        if any(k in h for k in keywords):
            return i

    raise ColumnNotFound(keywords[0], headers)


def load_font(size):
    for candidate in ("arial.ttf", "Arial.ttf", "calibri.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def load_as_image(src_path):
    """Open any supported file as a Pillow RGBA image.

    PDFs are rendered (first page) at PDF_DPI so they're sharp enough to print.
    """
    if src_path.suffix.lower() in PDF_EXTS:
        if fitz is None:
            raise RuntimeError("PDF support needs PyMuPDF. Run: pip install pymupdf")
        doc = fitz.open(src_path)
        page = doc.load_page(0)  # passport scans are page 1
        pix = page.get_pixmap(dpi=PDF_DPI)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        doc.close()
        return img.convert("RGBA")
    return Image.open(src_path).convert("RGBA")


def stamp_image(src_path, out_path, text, style):
    img = load_as_image(src_path)
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    font = load_font(style["font_size"])

    # text may have several lines (e.g. the index number above the ID number)
    spacing = max(4, style["font_size"] // 5)
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    pad = style["box_padding"]
    mx, my = style["margin_x"], style["margin_y"]
    W, H = img.size

    pos = style["position"]
    if "left" in pos:
        x = mx + pad
    else:  # right / center default right
        x = W - tw - mx - pad
    if "top" in pos:
        y = my + pad
    else:
        y = H - th - my - pad

    if style["box"]:
        box_rgb = hex_to_rgb(style["box_color"]) + (style["box_opacity"],)
        draw.rectangle(
            [x - pad, y - pad, x + tw + pad, y + th + pad],
            fill=box_rgb,
        )

    draw.multiline_text((x - bbox[0], y - bbox[1]), text, font=font,
                        spacing=spacing,
                        fill=hex_to_rgb(style["font_color"]) + (255,))

    result = Image.alpha_composite(img, overlay)

    # jpg and pdf can't hold transparency -> flatten to RGB
    if out_path.suffix.lower() in (".jpg", ".jpeg", ".pdf"):
        result = result.convert("RGB")
    result.save(out_path)


def build_excel_report(cfg, image_keys, not_found=None, stats=None,
                       needs_manual=None):
    """Write report.xlsx:
      - main sheet: same columns as the list, green=received / red=missing
      - "Summary" sheet: the counts
      - "Needs manual check" sheet: uncertain items to verify by hand
      - "Photos not found" sheet: pictures matching nobody
    """
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment

    xlsx = find_excel(cfg["excel_file"])
    df = pd.read_excel(xlsx, dtype=str).fillna("")
    headers = list(df.columns)

    name_col = _pick_column(headers, cfg.get("name_column", "auto"),
                            ("name", "client", "nom", "prenom", "full"))
    try:
        surname_col = _pick_column(headers, cfg.get("surname_column", "auto"),
                                   ("surname", "lastname", "family"))
    except ColumnNotFound:
        surname_col = None

    def clean_cell(v):
        v = str(v).strip()
        if v.endswith(".0") and v[:-2].isdigit():   # 16937265.0 -> 16937265
            v = v[:-2]
        return "" if v.lower() == "nan" else v

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    green_font = Font(color="006100")
    red_font = Font(color="9C0006")
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="404040")

    wb = Workbook()
    ws = wb.active
    ws.title = "Passport check"

    # pandas names blank columns "Unnamed: N" - show them empty instead
    clean_headers = ["" if str(h).startswith("Unnamed") else str(h)
                     for h in headers]
    out_headers = clean_headers + ["Photo?"]
    ws.append(out_headers)
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    received = missing = 0
    for _, row in df.iterrows():
        values = [clean_cell(row[h]) for h in headers]
        name = values[name_col] if name_col < len(values) else ""
        surname = (values[surname_col]
                   if surname_col is not None and surname_col < len(values) else "")
        key = normalize(f"{name} {surname}")
        found = bool(key) and key in image_keys

        values.append("YES - photo here" if found else "NO - no photo")
        ws.append(values)

        r = ws.max_row
        fill, font = (green, green_font) if found else (red, red_font)
        for cell in ws[r]:
            cell.fill = fill
            cell.font = font
        received += found
        missing += not found

    # column widths
    for i, h in enumerate(out_headers, start=1):
        longest = max([len(str(h))] +
                      [len(str(ws.cell(row=r, column=i).value or ""))
                       for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = \
            min(max(longest + 2, 8), 40)

    ws.freeze_panes = "A2"

    big = Font(bold=True, size=13, color="1F2937")

    # SUMMARY sheet (put first)
    if stats:
        wsum = wb.create_sheet("Summary", 0)
        wsum["A1"] = "PASSPORT CHECK - SHORT SUMMARY"
        wsum["A1"].font = big
        wsum["A2"] = (f"App version {__version__}  -  "
                      f"{datetime.datetime.now():%Y-%m-%d %H:%M}")
        wsum["A2"].font = Font(color="6B7280", size=9)
        line = [
            ("People in the list / Personnes dans la liste",
             stats["total"], None),
            ("Photos we have / Photos reçues",
             f'{stats["received"]}  ({stats["percent"]}%)', "C6EFCE"),
            ("People with no photo yet / Sans photo",
             stats["missing"], "FFC7CE"),
            ("Photos that match no name / Photos sans nom trouvé",
             stats["photos_not_found"], "FFEB9C"),
            ("Photos with the number added / Numéro ajouté",
             stats["stamped"], "C6EFCE"),
            ("Photos with list number only / Numéro de liste seul",
             stats.get("index_only", 0), "DBEAFE"),
            ("Problem files / Fichiers à problème",
             stats.get("broken", 0), "FFC7CE"),
            ("PDF files still to sort / PDF à trier",
             stats["pdf"], None),
            ("PLEASE CHECK BY HAND / À VÉRIFIER À LA MAIN",
             stats["needs_manual"], "E9D5FF"),
        ]
        r = 3
        for label, value, color in line:
            wsum.cell(r, 1, label).font = Font(bold=(label.isupper()))
            c = wsum.cell(r, 2, value)
            if color:
                fill = PatternFill("solid", fgColor=color)
                wsum.cell(r, 1).fill = fill
                c.fill = fill
            r += 1
        wsum.column_dimensions["A"].width = 52
        wsum.column_dimensions["B"].width = 18

    # NEEDS MANUAL CHECK sheet (highlighted - verify to be 100% sure)
    wsm = wb.create_sheet("Please check by hand")
    wsm.append(["Photo or name / Photo ou nom",
                "What to do (English)", "Quoi faire (Français)",
                "New file name / Nouveau nom"])
    for cell in wsm[1]:
        cell.font = head_font
        cell.fill = head_fill
    purple_fill = PatternFill("solid", fgColor="E9D5FF")
    purple_font = Font(color="6B21A8")
    for item, why, why_fr, sugg in (needs_manual or []):
        wsm.append([item, why, why_fr, sugg])
        for cell in wsm[wsm.max_row]:
            cell.fill = purple_fill
            cell.font = purple_font
    for col, w in (("A", 36), ("B", 60), ("C", 60), ("D", 32)):
        wsm.column_dimensions[col].width = w
    wsm.freeze_panes = "A2"

    # photos we could not match to anyone in the list - with the exact reason
    ws2 = wb.create_sheet("Photos with no match")
    ws2.append(["Photo file / Fichier",
                "Why / Pourquoi",
                "What to do (English)", "Quoi faire (Français)",
                "New file name / Nouveau nom"])
    for cell in ws2[1]:
        cell.font = head_font
        cell.fill = head_fill
    warn_fill = PatternFill("solid", fgColor="FFEB9C")   # amber
    warn_font = Font(color="9C6500")
    reason_words = {"reversed": "name back-to-front / nom à l'envers",
                    "junk": "extra words / mots en trop",
                    "two_people": "two people in one file / deux personnes",
                    "spelling": "spelling mistake / faute de frappe",
                    "partial": "only part matches / correspondance partielle",
                    "unknown": "not in the list / pas dans la liste"}
    for fname, diag in (not_found or []):
        ws2.append([fname, reason_words.get(diag["reason"], diag["reason"]),
                    diag["why"], diag["why_fr"], diag["suggestion"]])
        for cell in ws2[ws2.max_row]:
            cell.fill = warn_fill
            cell.font = warn_font
    for col, w in (("A", 34), ("B", 34), ("C", 55), ("D", 55), ("E", 32)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    report_x = HERE / "report.xlsx"
    try:
        wb.save(report_x)
    except PermissionError:
        # file is open in Excel -> save under a fallback name instead of crashing
        report_x = HERE / "report (new).xlsx"
        wb.save(report_x)
        print("  NOTE: report.xlsx was open in Excel. Close it next time.")
    print(f"  Excel report: {report_x.name}  "
          f"(green received: {received}, red missing: {missing}, "
          f"photos not found: {len(not_found or [])})")
    return report_x


# --------------------------------------------------------------------------- #
# core - callable by both the CLI and the app
# --------------------------------------------------------------------------- #
def run_check(cfg=None, log=print, progress=None):
    """Do the whole job. Returns a results dict the UI can display.

    progress: optional callback(i, total, filename) called for every photo.
    """
    if cfg is None:
        cfg = load_config()
    style = cfg["id_text"]
    images_dir = HERE / cfg["images_folder"]
    out_dir = HERE / cfg["output_folder"]
    out_dir.mkdir(exist_ok=True)

    log(f"Passport Check v{__version__}")
    log("Reading Excel list...")
    people, duplicates = read_excel(cfg)
    log(f"  {len(people)} people with a unique name.")
    if duplicates:
        log(f"  {len(duplicates)} name(s) appear more than once "
            f"-> flagged for manual handling.")
    log("")

    skip_pdf = cfg.get("skip_pdf", True)

    # accept known image types, PDFs, and extension-less files (some scans have none)
    def is_candidate(p):
        if not p.is_file() or p.name.startswith(("~$", ".")):
            return False
        suf = p.suffix.lower()
        return suf in ALL_EXTS or suf == ""

    images = [p for p in images_dir.iterdir() if is_candidate(p)]

    matched, unknown_images, no_id, manual, pdf_skipped = [], [], [], [], []
    index_only = []   # photo stamped with the index only (no passport number yet)
    broken = []       # (filename, plain-words problem) - unreadable files etc.
    seen_names = set()
    used_photo = {}   # key -> first filename stamped (to catch double photos)
    decisions = []    # one entry per photo: what we decided and why (audit log)

    sorted_images = sorted(images)
    total_imgs = len(sorted_images)

    for i, img in enumerate(sorted_images, start=1):
        if progress:
            progress(i, total_imgs, img.name)

        # PDFs: skip for now (they may contain several passports on one file)
        if img.suffix.lower() in PDF_EXTS and skip_pdf:
            pdf_skipped.append(img.name)
            decisions.append({"file": img.name, "key": "",
                              "decision": "pdf skipped",
                              "detail": "PDFs are sorted separately"})
            continue

        key = clean_stem(img.stem)
        if key in duplicates:
            # same name exists several times with different IDs -> do NOT guess
            seen_names.add(key)
            manual.append((img.name, duplicates[key]["ids"]))
            decisions.append({"file": img.name, "key": key,
                              "decision": "duplicate name - check by hand",
                              "matched_name": duplicates[key]["name"],
                              "detail": "possible numbers: "
                                        + ", ".join(duplicates[key]["ids"])})
            continue
        if key in people:
            person = people[key]
            seen_names.add(key)

            # two photo files for the same person (e.g. JOHN.jpg + JOHN.png)?
            if key in used_photo:
                broken.append((img.name,
                               f"There are two photos for this person. "
                               f"We used '{used_photo[key]}'. Delete this "
                               f"one if it is a double.",
                               f"Il y a deux photos pour cette personne. "
                               f"Nous avons utilisé '{used_photo[key]}'. "
                               f"Supprimez celle-ci si c'est un doublon."))
                decisions.append({"file": img.name, "key": key,
                                  "decision": "double photo - not used",
                                  "matched_name": person["name"],
                                  "detail": f"already used "
                                            f"'{used_photo[key]}'"})
                continue

            idx = str(person.get("index", "")).strip()
            has_id = bool(person["id"])

            # build the text. Normally: index line + ID line.
            # If there is NO passport number, fall back to the index only,
            # so every photo still gets a printable marker.
            lines = []
            if (style.get("show_index", True) or not has_id) and idx:
                lines.append(f"{style.get('index_prefix', 'No ')}{idx}")
            if has_id:
                lines.append(f"{style['prefix']}{person['id']}")
            if not lines:
                no_id.append(person["name"])   # nothing at all to write
                decisions.append({"file": img.name, "key": key,
                                  "decision": "nothing to write",
                                  "matched_name": person["name"],
                                  "detail": "no passport number and no "
                                            "list position in the Excel"})
                continue
            text = "\n".join(lines)

            # output filename: number it by Excel index so prints come out in order
            suffix = img.suffix or ".jpg"
            if cfg.get("number_output_files", True) and idx.isdigit():
                out_name = f"{int(idx):03d} - {img.stem}{suffix}"
            else:
                out_name = img.stem + suffix
            out_path = out_dir / out_name

            # one bad file must never stop the whole run
            try:
                if img.stat().st_size == 0:
                    raise UnidentifiedImageError("empty file")
                stamp_image(img, out_path, text, style)
            except (UnidentifiedImageError, OSError, RuntimeError,
                    ValueError) as e:
                if isinstance(e, PermissionError):
                    msg = ("The file is open in another program. "
                           "Close it and run again.")
                    msg_fr = ("Le fichier est ouvert dans un autre "
                              "programme. Fermez-le et relancez.")
                elif img.suffix.lower() in PDF_EXTS:
                    msg = ("This PDF cannot be opened (it may be damaged "
                           "or have a password). Ask for the file again.")
                    msg_fr = ("Ce PDF ne peut pas être ouvert (peut-être "
                              "endommagé ou protégé par mot de passe). "
                              "Redemandez le fichier.")
                else:
                    msg = ("This file is damaged or is not a real photo. "
                           "Ask the client to send it again.")
                    msg_fr = ("Ce fichier est endommagé ou n'est pas une "
                              "vraie photo. Demandez au client de la "
                              "renvoyer.")
                log_error(f"stamping {img.name}", e)
                broken.append((img.name, msg, msg_fr))
                decisions.append({"file": img.name, "key": key,
                                  "decision": "problem file - not used",
                                  "matched_name": person["name"],
                                  "detail": msg})
                continue

            used_photo[key] = img.name
            if has_id:
                matched.append((idx, person["name"], person["id"]))
                decisions.append({"file": img.name, "key": key,
                                  "decision": "matched - number added",
                                  "matched_name": person["name"],
                                  "id": person["id"], "index": idx,
                                  "detail": "exact name match"})
            else:
                index_only.append((idx, person["name"]))
                decisions.append({"file": img.name, "key": key,
                                  "decision": "matched - list number only",
                                  "matched_name": person["name"],
                                  "index": idx,
                                  "detail": "exact name match, but no "
                                            "passport number in the Excel"})
        else:
            unknown_images.append((img.name, key))

    missing_image = [p["name"] for k, p in people.items() if k not in seen_names]

    # every real photo we received (matched OR not), for the coloured Excel
    image_keys = {clean_stem(p.stem) for p in images
                  if p.suffix.lower() not in PDF_EXTS}

    # for each not-found photo, work out exactly WHY and what to do
    not_found = []   # (filename, diag dict)
    for fname, key in unknown_images:
        diag = diagnose_unmatched(key, fname, people, duplicates)
        not_found.append((fname, diag))
        decisions.append({"file": fname, "key": key,
                          "decision": f"no match ({diag['reason']})",
                          "detail": diag["why"],
                          "suggestion": diag["suggestion"]})
    for name in missing_image:
        decisions.append({"file": "", "key": normalize(name),
                          "decision": "no photo received",
                          "matched_name": name,
                          "detail": "in the Excel list but no photo file"})

    # ------------------------------------------------------------------ report
    log("=" * 60)
    log("RESULTS")
    log("=" * 60)
    log(f"\nNumber added to {len(matched)} photo(s) -> 'output' folder:")
    for idx, name, cid in matched:
        log(f"     - {name}  (number {cid or '??? MISSING'})")

    if index_only:
        log(f"\nList number added (no passport number yet) "
            f"({len(index_only)}):")
        for idx, name in index_only:
            log(f"     - {name}  (No {idx})")

    if manual:
        log(f"\nCHECK BY HAND - same name twice ({len(manual)}):")
        for fname, ids in manual:
            log(f"     - {fname}   numbers: {', '.join(ids)}")

    if no_id:
        log(f"\nPhoto here but no number and no list position ({len(no_id)}):")
        for n in no_id:
            log(f"     - {n}")

    if pdf_skipped:
        log(f"\nPDF files still to sort ({len(pdf_skipped)}):")
        for n in pdf_skipped:
            log(f"     - {n}")

    if broken:
        log(f"\nPROBLEM FILES - could not be used ({len(broken)}):")
        for fname, msg, _fr in broken:
            log(f"     - {fname}   -> {msg}")

    if missing_image:
        log(f"\nPeople with no photo yet ({len(missing_image)}):")
        for n in missing_image:
            log(f"     - {n}")

    if not_found:
        log(f"\nPhotos that match no name in the list ({len(not_found)}):")
        for fname, diag in not_found:
            log(f"     - {fname}")
            log(f"         {diag['why']}")

    # ---- stats + the "verify manually to be 100% sure" list ----------------
    total = len(people) + len(duplicates)
    received = len({k for k in ({**people, **duplicates}) if k in image_keys})
    pct = round(received / total * 100) if total else 0

    # (item, why, why_fr, suggested_new_filename)
    needs_manual = []
    for fname, diag in not_found:
        needs_manual.append((fname, diag["why"], diag["why_fr"],
                             diag["suggestion"]))
    for fname, ids in manual:   # same name twice, different IDs
        needs_manual.append(
            (fname,
             f"Two people have this same name. Choose the right "
             f"number: {', '.join(ids)}",
             f"Deux personnes ont ce même nom. Choisissez le bon "
             f"numéro : {', '.join(ids)}",
             ""))
    for fname, msg, msg_fr in broken:
        needs_manual.append((fname, msg, msg_fr, ""))

    stats = {
        "total": total, "received": received, "missing": total - received,
        "percent": pct, "stamped": len(matched),
        "index_only": len(index_only),
        "photos_not_found": len(not_found),
        "no_id": len(no_id), "pdf": len(pdf_skipped),
        "broken": len(broken),
        "needs_manual": len(needs_manual),
    }

    log("\n" + "-" * 60)
    log(f"SUMMARY: we have {received} of {total} photos ({pct}%). "
        f"Number added to {len(matched)}, list number only to {len(index_only)}. "
        f"Please check {len(needs_manual)} by hand.")

    # save the text report
    report = HERE / "report.txt"
    with open(report, "w", encoding="utf-8") as f:
        f.write("PASSPORT CHECK - REPORT\n")
        f.write(f"App version {__version__}  -  "
                f"{datetime.datetime.now():%Y-%m-%d %H:%M}\n")
        f.write("=" * 40 + "\n")
        f.write("SHORT SUMMARY / RÉSUMÉ\n")
        f.write(f"  People in the list              : {total}\n")
        f.write(f"  Photos we have                  : {received}  ({pct}%)\n")
        f.write(f"  People with no photo yet        : {total - received}\n")
        f.write(f"  Photos that match no name       : {len(not_found)}\n")
        f.write(f"  Photos with the number added    : {len(matched)}\n")
        f.write(f"  Photos with list number only    : {len(index_only)}\n")
        f.write(f"  Problem files (damaged...)      : {len(broken)}\n")
        f.write(f"  PLEASE CHECK BY HAND            : {len(needs_manual)}\n")
        f.write("=" * 40 + "\n\n")

        if needs_manual:
            f.write(f">>> PLEASE CHECK THESE BY HAND / À VÉRIFIER À LA MAIN "
                    f"({len(needs_manual)}):\n")
            for item, why, why_fr, sugg in needs_manual:
                f.write(f"  ! {item}\n      {why}\n      {why_fr}\n")
                if sugg:
                    f.write(f"      -> new file name / nouveau nom : {sugg}\n")
            f.write("\n")

        f.write(f"Number added to these photos ({len(matched)}):\n")
        for idx, name, cid in matched:
            f.write(f"  {name}\t{cid}\n")
        f.write(f"\nPeople with no photo yet ({len(missing_image)}):\n")
        for n in missing_image:
            f.write(f"  {n}\n")
        f.write(f"\nPhotos that match no name in the list "
                f"({len(not_found)}):\n")
        for fname, diag in not_found:
            f.write(f"  {fname}\n      {diag['why']}\n")
        f.write(f"\nSame name twice - choose the number ({len(manual)}):\n")
        for fname, ids in manual:
            f.write(f"  {fname}\tnumbers: {', '.join(ids)}\n")
        f.write(f"\nProblem files ({len(broken)}):\n")
        for fname, msg, _fr in broken:
            f.write(f"  {fname}\t{msg}\n")
        f.write(f"\nList number added, still waiting for passport number "
                f"({len(index_only)}):\n")
        for idx, name in index_only:
            f.write(f"  No {idx}\t{name}\n")
        f.write(f"\nPDF files still to sort ({len(pdf_skipped)}):\n")
        for n in pdf_skipped:
            f.write(f"  {n}\n")
        if no_id:
            f.write(f"\nPhoto here but no number and no list position ({len(no_id)}):\n")
            for n in no_id:
                f.write(f"  {n}\n")

    log(f"\nText report saved to: {report.name}")

    # coloured Excel report (summary sheet + not-found + needs-manual sheets)
    report_x = build_excel_report(cfg, image_keys, not_found, stats,
                                  needs_manual)

    # decision log next to the reports (why every photo went where it went)
    decisions_csv = HERE / "decisions.csv"
    try:
        write_decisions_csv(decisions_csv, decisions)
        log(f"Decision log saved to: {decisions_csv.name}")
    except OSError as e:
        log_error("writing decisions.csv", e)
        decisions_csv = None
        log("NOTE: decisions.csv is open in another program - "
            "the decision log was not saved this time.")

    # keep a dated copy of everything in runs/ so testing runs can be compared
    run_dir = archive_run(decisions, stats, cfg,
                          [report, report_x, decisions_csv])
    if run_dir:
        log(f"Run archived in: runs\\{run_dir.name}")
    log("Done.")

    return {
        "stamped": matched,               # (index, name, id)
        "index_only": index_only,         # (index, name) - no passport number yet
        "no_id": no_id,
        "pdf_skipped": pdf_skipped,
        "missing_image": missing_image,
        "not_found": not_found,           # (filename, diag dict)
        "manual": manual,
        "broken": broken,                 # (filename, why, why_fr)
        "needs_manual": needs_manual,     # (item, why, why_fr, suggestion)
        "stats": stats,
        "total_people": len(people) + len(duplicates),
    }


def main():
    run_check(cfg=None, log=print)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n[ERROR] {e}")
    input("\nPress Enter to close...")
